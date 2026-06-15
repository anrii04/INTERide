"""
persistence.py
INTERide — File I/O: JSON booking persistence + .docx receipt export.
"""

from typing import List, Tuple
import json
import os

from models import Booking

FILE: str = "interide_bookings.json"


# ══════════════════════════════════════════════════════════════════════════
#  JSON persistence
# ══════════════════════════════════════════════════════════════════════════
def save_all(blist: List[Booking]) -> None:
    """Serialise every booking to *FILE*."""
    with open(FILE, "w") as f:
        json.dump([b.to_dict() for b in blist], f, indent=2)


def load_all() -> List[Booking]:
    """Deserialise bookings from *FILE*; returns [] on missing / corrupt file."""
    if not os.path.exists(FILE):
        return []
    try:
        with open(FILE) as f:
            return [Booking.from_dict(d) for d in json.load(f)]
    except Exception:
        return []


# convenience aliases used throughout the screens
save_bookings = save_all
load_bookings = load_all


# ══════════════════════════════════════════════════════════════════════════
#  DOCX receipt
# ══════════════════════════════════════════════════════════════════════════
def make_receipt(b: Booking, path: str) -> None:
    """Generate a formatted Word receipt for booking *b* and save to *path*."""
    from config import HAS_DOCX
    if not HAS_DOCX:
        raise RuntimeError("python-docx is not installed.")

    from docx import Document as DocxDoc
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = DocxDoc()
    for sec in doc.sections:
        sec.top_margin    = sec.bottom_margin = Cm(2)
        sec.left_margin   = sec.right_margin  = Cm(2.5)

    # ── Title ──
    p  = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r  = p.add_run("INTERide")
    r.bold = True
    r.font.size      = Pt(30)
    r.font.color.rgb = RGBColor(37, 99, 235)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Official Booking Receipt")
    r2.font.size      = Pt(11)
    r2.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph()

    # ── Data table ──
    rows: List[Tuple[str, str]] = [
        ("Booking ID",  b.bid),
        ("Passenger",   b.user),
        ("Date & Time", b.ts),
        ("Vehicle",     f"{b.emoji} {b.vtype}"),
        ("Pickup",      b.pick),
        ("Drop-off",    b.drop),
        ("Distance",    f"{b.km} km"),
        ("Total Fare",  f"₱{b.cost:,.2f}"),
    ]
    tbl = doc.add_table(rows=len(rows), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style     = "Table Grid"

    for i, (lbl, val) in enumerate(rows):
        rc = tbl.rows[i].cells
        lr = rc[0].paragraphs[0].add_run(lbl)
        lr.bold            = True
        lr.font.size       = Pt(10)
        lr.font.color.rgb  = RGBColor(80, 80, 80)
        vr = rc[1].paragraphs[0].add_run(val)
        vr.font.size       = Pt(10)
        if lbl == "Total Fare":
            vr.bold            = True
            vr.font.color.rgb  = RGBColor(37, 99, 235)
            vr.font.size       = Pt(13)

    # ── Footer ──
    doc.add_paragraph()
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Thank you for riding with INTERide! Safe travels.")
    fr.font.size      = Pt(9)
    fr.font.color.rgb = RGBColor(150, 150, 150)
    fr.font.italic    = True

    doc.save(path)
